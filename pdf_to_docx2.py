import uuid
import fitz
from pdf2docx import Converter
from docx import Document


def merge_docx(docx_list, output_path):
    merged_document = Document(docx_list[0])

    for docx_path in docx_list[1:]:
        sub_doc = Document(docx_path)

        for paragraph in sub_doc.paragraphs:
            new_p = merged_document.add_paragraph()
            if paragraph.style:
                new_p.style = paragraph.style
            for run in paragraph.runs:
                new_run = new_p.add_run(run.text)
                new_run.bold = run.bold
                new_run.italic = run.italic
                new_run.underline = run.underline

        for table in sub_doc.tables:
            new_table = merged_document.add_table(rows=0, cols=len(table.columns))
            for row in table.rows:
                new_row = new_table.add_row().cells
                for idx, cell in enumerate(row.cells):
                    new_row[idx].text = cell.text

        merged_document.add_page_break()

    merged_document.save(output_path)


def split_pdf_docx(pdf, output_docx):
    cv = Converter(pdf)
    pages = []
    page_count = get_pdf_page_count(pdf)

    for page in range(page_count):
        _uuid = uuid.uuid4()
        file_name = f"./file/temp/{page}_{_uuid}.docx"
        cv.convert(file_name, start=page, end=page + 1)
        pages.append(file_name)

    cv.close()

    # 合并成一个 docx
    merge_docx(pages, output_docx)

    print(f"共生成 {len(pages)} 份 docx，已合并到：{output_docx}")


def get_pdf_page_count(pdf_path):
    doc = fitz.open(pdf_path)
    return doc.page_count


if __name__ == '__main__':
    pdf = "./file/传统医学师承关系合同书.pdf"
    docx = "./file/传统医学师承关系合同书.docx"
    split_pdf_docx(pdf, docx)
