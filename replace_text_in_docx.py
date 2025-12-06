import re

from docx import Document


def replace_text_in_docx(input_path, output_path, replace_map):
    """
    replace_map: dict，例如 {"旧文本": "新文本", "姓名": "张三"}
    """
    doc = Document(input_path)

    # 处理段落
    for paragraph in doc.paragraphs:
        replace_runs_in_paragraph(paragraph, replace_map)

    # 处理表格（docx 中很多内容在表格里）
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_runs_in_paragraph(paragraph, replace_map)

    doc.save(output_path)


def replace_runs_in_paragraph(paragraph, replace_map):
    """处理 paragraph.runs，实现跨 run 替换"""
    if not paragraph.runs:
        return

    # 把所有 runs 的文字合并
    full_text = "".join(run.text for run in paragraph.runs)

    # 逐个执行替换
    for old, new in replace_map.items():
        full_text = re.sub(re.escape(old), new, full_text)

    # 清空原 runs
    for run in paragraph.runs:
        run.text = ""

    # 放回新的内容到第一个 run
    paragraph.runs[0].text = full_text


def count_paragraphs_and_table_paragraphs(docx_path):
    doc = Document(docx_path)

    # 文档中的普通段落数量
    paragraph_count = len(doc.paragraphs)

    # 表格中的单元格段落数量
    table_paragraph_count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                # 每个 cell 内可能包含多个段落
                table_paragraph_count += len(cell.paragraphs)

    total = paragraph_count + table_paragraph_count

    return paragraph_count, table_paragraph_count, total


if __name__ == '__main__':
    replace_map = {
        "NAME": "张三",
        "DATE": "2025-12-06",
        "项目名称": "智能AI助手"
    }
    replace_text_in_docx("./file/input.docx", "./file/output.docx", replace_map)
    print("转换成功")

    p_count, t_count, total = count_paragraphs_and_table_paragraphs("./file/input.docx")
    print("普通段落数量:", p_count)
    print("表格段落数量:", t_count)
    print("总段落数量:", total)
